"""Maakt tests/fixture.xlsx: een nep-actieregister met dezelfde opbouw als het echte
bestand (titelregels, Excel-tabel, formules in Achterstallig/Dagen open, dropdowns).
Bevat uitsluitend verzonnen data.  Gebruik: python3 tests/make_fixture.py
"""
import datetime as dt
import zipfile
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

HEADERS = ["Nr", "Actie", "Domein", "Eigenaar", "Prioriteit", "Deadline", "Status",
           "Achterstallig", "Dagen open", "Rang", "Waarom nu", "Aanpak — eerste stap",
           "Bron", "Besluitdatum", "Toelichting"]
DOMEINEN = ["A. Strategie & Beleid", "B. Personeel en HRM", "D. ICT en digitalisering"]
STATUSSEN = ["Idee", "Plan", "Open", "Loopt", "Wacht op", "Gepland", "Te verifiëren", "Klaar"]

ROWS = [
    (1, "Testactie klaar", DOMEINEN[1], "Anna", "P1", dt.date(2026, 8, 31), "Klaar", 3,
     "Omdat het moet.", "Eerst bellen.", "Verslag 1", dt.date(2026, 7, 8), "Afgerond."),
    (2, "Testactie achterstallig", DOMEINEN[2], "Anna / Bert", "P1", dt.date(2026, 9, 1), "Open", 1,
     "Blokkeert anderen.", "Mail sturen.", "Verslag 2", dt.date(2026, 8, 12), ""),
    (3, "Testactie zonder deadline", DOMEINEN[0], "Bert", "P3", None, "Idee", None,
     "", "", "Braindump", None, "Nog besluiten."),
    (4, "Testactie wacht op", DOMEINEN[2], "Carla", "P2", dt.date(2026, 12, 31), "Wacht op", None,
     "Leverancier moet leveren.", "Rappelleren.", "Mail", dt.date(2026, 10, 1), ""),
]


def main():
    wb = Workbook()
    ws = wb.active
    ws.title = "Actieregister"
    ws["A1"] = "Actieregister — Testpraktijk"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = "Maatschapniveau · versie 01-10-2026 · testbestand met nepdata."
    for c, h in enumerate(HEADERS, 1):
        ws.cell(row=4, column=c, value=h)
    for i, r in enumerate(ROWS):
        row = 5 + i
        nr, actie, dom, eig, prio, deadline, status, rang, waarom, aanpak, bron, besluit, toel = r
        values = [nr, actie, dom, eig, prio, deadline, status, None, None, rang, waarom, aanpak,
                  bron, besluit, toel]
        for c, v in enumerate(values, 1):
            cell = ws.cell(row=row, column=c, value=v)
            if isinstance(v, dt.date):
                cell.number_format = "DD-MM-YYYY"
        ws.cell(row=row, column=8,
                value=f'=IF(AND(F{row}<>"",F{row}<TODAY(),G{row}<>"Klaar"),"JA","")')
        ws.cell(row=row, column=9,
                value=f'=IF(N{row}="","",IF(G{row}="Klaar","",TODAY()-N{row}))')
        ws.cell(row=row, column=1).fill = PatternFill("solid", fgColor="DDEEFF")
    last = 4 + len(ROWS)
    table = Table(displayName="Actieregister", ref=f"A4:O{last}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(table)

    dv = DataValidation(type="list", formula1='"' + ",".join(STATUSSEN) + '"', allow_blank=True)
    dv.add(f"G5:G{last}")
    ws.add_data_validation(dv)
    dv2 = DataValidation(type="list", formula1='"P1,P2,P3"', allow_blank=True)
    dv2.add(f"E5:E{last}")
    ws.add_data_validation(dv2)

    ov = wb.create_sheet("Overzicht")
    ov["A1"] = "Overzicht"
    ov["A2"] = "Open acties totaal"
    ov["B2"] = '=COUNTIFS(Actieregister[Status],"<>Klaar",Actieregister[Status],"<>Idee")'
    tl = wb.create_sheet("Toelichting")
    tl["A1"] = "Zo werkt dit register"

    out = Path(__file__).with_name("fixture.xlsx")
    wb.save(out)
    excel_style_rels(out)
    print(f"geschreven: {out}")


def excel_style_rels(path):
    """openpyxl schrijft absolute paden ("/xl/tables/table1.xml") in de relaties;
    Excel zelf schrijft relatieve ("../tables/table1.xml"). Zet ze om zoals Excel."""
    with zipfile.ZipFile(path) as z:
        items = [(i, z.read(i.filename)) for i in z.infolist()]
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        for info, data in items:
            if info.filename.startswith("xl/worksheets/_rels/"):
                data = data.replace(b'Target="/xl/tables/', b'Target="../tables/')
            z.writestr(info, data)


if __name__ == "__main__":
    main()
